"""Generate a self-contained HTML coverage map: Excel  <->  v3 UI data contract.

Reads `v3/assets/Data from Miro.xlsx` and `server/app/data/fallback.json` (current v3
structure) and emits `excel-mapping.html` at the repo root showing:
  - source summary (sheets, rows)
  - phase/subphase structure: v3 (current) vs Excel (incoming)
  - field-by-field coverage (mapped / from config / dropped-deferred)
  - the 84-item enabler catalog grouped by phase+type (to be inserted within subphases)

Run:  cd server && uv run python scripts/build_excel_mapping.py
"""

from __future__ import annotations

import html
import json
from collections import Counter, defaultdict
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[2]
XLSX = ROOT / "v3" / "assets" / "Data from Miro.xlsx"
FALLBACK = ROOT / "server" / "app" / "data" / "fallback.json"
OUT = ROOT / "excel-mapping.html"


def esc(v) -> str:
    return html.escape("" if v is None else str(v))


# ── Read Excel ────────────────────────────────────────────────────────────────
wb = openpyxl.load_workbook(XLSX, data_only=True)
jtbd_ws = wb["JTBD"]
enabler_ws = wb["Overview MFG enablers "]

jrows = list(jtbd_ws.iter_rows(values_only=True))[2:]  # skip header + Now/Near/Future row
excel = defaultdict(lambda: defaultdict(list))  # phase -> subphase -> [(no,title)]
phase = sub = None
for r in jrows:
    if r[0]:
        phase = r[0]
    if r[1]:
        sub = r[1]
    if r[4] or r[5]:
        excel[phase][sub].append((r[4], r[5]))

enabler_rows = [r for r in list(enabler_ws.iter_rows(values_only=True))[1:] if r[1]]
by_phase_type = defaultdict(lambda: defaultdict(list))
for r in enabler_rows:
    by_phase_type[r[4]][r[2]].append(r[1])

# ── Read current v3 ───────────────────────────────────────────────────────────
fb = json.loads(FALLBACK.read_text())
v3 = [(p["title"], [s["title"] for s in p["subphases"]]) for p in fb["strategy"]["phases"]]

# ── Coverage rows: (UI field, Excel source, status) ──────────────────────────
COV = {
    "Phase": [
        ("title", "Phase column", "mapped"),
        ("id", "slug(Phase)", "derived"),
        ("subtitle", "—", "config"),
        ("accent (colour)", "—", "config"),
        ("summary", "—", "config / dropped"),
        ("signal", "—", "dropped"),
    ],
    "Subphase": [
        ("title", "Milestone", "mapped"),
        ("summary", "Milestone description", "mapped"),
        ("id", "slug(Milestone)", "derived"),
        ("milestone (statement)", "—", "dropped"),
        ("enablers (in subphase)", "Sheet 2 catalog (by phase/span)", "mapped*"),
    ],
    "JTBD": [
        ("number", "JTBD No", "mapped"),
        ("title", "JTBD", "mapped"),
        ("summary", "Description", "mapped"),
        ("status", "set ‘Sample complete’", "derived"),
        ("owner / effort", "—", "dropped"),
        ("sections.coreActivities", "Core activities", "mapped"),
        ("sections.inputs", "Inputs", "mapped"),
        ("sections.outputs", "Outputs", "mapped"),
        ("sections.enablers.now/near/next.tools", "Enablers (tools) Now/Near/Future", "mapped"),
        ("sections.enablers.now/near/next.agents", "Enablers (agents) Now/Near/Future", "mapped"),
        ("sections.valueDriver", "—", "dropped (Figma)"),
        ("sections.coreProblem", "—", "dropped (Figma)"),
        ("(Training column)", "Training", "unused"),
    ],
    "Lanes": [
        ("jtbd", "JTBD rows", "mapped"),
        ("opportunities", "Key opportunities blob → mfgFocus", "mapped"),
        ("enablers", "Sheet 2 catalog", "mapped*"),
        ("enabledBy", "—", "dropped (Figma)"),
        ("valueCreation", "—", "dropped (Figma)"),
    ],
    "Views": [
        ("Value Tree", "—", "dropped (Figma)"),
        ("Timeline (dated)", "Sheet 2 has no dates", "dropped → replaced by catalog"),
    ],
}

STATUS_CLASS = {
    "mapped": "ok",
    "mapped*": "ok",
    "derived": "ok",
    "config": "warn",
    "config / dropped": "warn",
    "unused": "muted",
    "dropped": "drop",
    "dropped (Figma)": "drop",
    "dropped → replaced by catalog": "drop",
}


def cov_table(title, rows):
    out = [f'<h3>{esc(title)}</h3><table><thead><tr><th>UI field</th><th>Excel source</th><th>Status</th></tr></thead><tbody>']
    for field, src, status in rows:
        cls = STATUS_CLASS.get(status, "muted")
        out.append(
            f'<tr><td><code>{esc(field)}</code></td><td>{esc(src)}</td>'
            f'<td><span class="pill {cls}">{esc(status)}</span></td></tr>'
        )
    out.append("</tbody></table>")
    return "".join(out)


# ── Structure comparison ──────────────────────────────────────────────────────
def struct_block():
    out = ['<div class="cols">']
    out.append('<div><h3>Current v3 (placeholders)</h3>')
    for ptitle, subs in v3:
        out.append(f'<div class="phase"><strong>{esc(ptitle)}</strong><ul>')
        for s in subs:
            out.append(f"<li>{esc(s)}</li>")
        out.append("</ul></div>")
    out.append("</div>")

    out.append('<div><h3>Incoming Excel (source of truth)</h3>')
    for ptitle, submap in excel.items():
        njtbd = sum(len(v) for v in submap.values())
        out.append(f'<div class="phase"><strong>{esc(ptitle)}</strong> <span class="muted">· {njtbd} JTBD</span><ul>')
        for s, items in submap.items():
            nums = ", ".join(esc(n) for n, _ in items)
            out.append(f"<li>{esc(s)} <span class='muted'>({nums})</span></li>")
        out.append("</ul></div>")
    out.append("</div></div>")
    return "".join(out)


def enabler_block():
    out = ['<table><thead><tr><th>Primary Phase</th><th>Tools</th><th>Agents</th><th>Total</th></tr></thead><tbody>']
    for ph, types in by_phase_type.items():
        tools = len(types.get("Tool", []))
        agents = len(types.get("Agent", [])) + len(types.get("Agent (Super)", []))
        out.append(f"<tr><td>{esc(ph)}</td><td>{tools}</td><td>{agents}</td><td>{tools+agents}</td></tr>")
    out.append(f"<tr class='tot'><td>Total</td><td>{sum(1 for r in enabler_rows if r[2]=='Tool')}</td>"
               f"<td>{sum(1 for r in enabler_rows if 'Agent' in (r[2] or ''))}</td><td>{len(enabler_rows)}</td></tr>")
    out.append("</tbody></table>")
    return "".join(out)


coverage_html = "".join(cov_table(k, v) for k, v in COV.items())

excel_phases = list(excel.keys())
v3_phases = [p for p, _ in v3]

HTML = f"""<!doctype html>
<html lang="en"><head><meta charset="utf-8">
<meta name="viewport" content="width=device-width, initial-scale=1">
<title>Excel → UI coverage map</title>
<link href="https://fonts.googleapis.com/css2?family=Google+Sans+Flex:wght@400;500;600;700;800&display=swap" rel="stylesheet">
<style>
:root{{--ink:#14171c;--muted:#6b7280;--paper:#f6f5f1;--line:rgba(20,23,28,.12);
--blue:#4285f4;--red:#ea4335;--green:#34a853;--yellow:#f9ab00;}}
*{{box-sizing:border-box}}
body{{margin:0;background:var(--paper);color:var(--ink);font-family:"Google Sans Flex",system-ui,sans-serif;
line-height:1.5;font-size:14px}}
.wrap{{max-width:1080px;margin:0 auto;padding:48px 28px 96px}}
h1{{font-size:30px;margin:0 0 6px;font-weight:700;letter-spacing:-.02em}}
.sub{{color:var(--muted);margin:0 0 32px;font-size:15px}}
h2{{font-size:13px;text-transform:uppercase;letter-spacing:.08em;color:var(--muted);
margin:44px 0 14px;border-bottom:1px solid var(--line);padding-bottom:8px}}
h3{{font-size:15px;margin:22px 0 10px;font-weight:600}}
.cards{{display:grid;grid-template-columns:repeat(4,1fr);gap:12px}}
.card{{background:#fff;border:1px solid var(--line);border-radius:12px;padding:16px}}
.card b{{font-size:26px;font-weight:700;display:block}}
.card span{{color:var(--muted);font-size:12px}}
.cols{{display:grid;grid-template-columns:1fr 1fr;gap:24px}}
.phase{{background:#fff;border:1px solid var(--line);border-radius:10px;padding:12px 14px;margin-bottom:10px}}
.phase ul{{margin:6px 0 0;padding-left:18px}} .phase li{{margin:2px 0}}
table{{width:100%;border-collapse:collapse;background:#fff;border:1px solid var(--line);
border-radius:10px;overflow:hidden;margin:6px 0 18px}}
th,td{{text-align:left;padding:9px 12px;border-bottom:1px solid var(--line);vertical-align:top}}
th{{background:#fafafa;font-weight:600;font-size:12px;text-transform:uppercase;letter-spacing:.04em;color:var(--muted)}}
tr:last-child td{{border-bottom:none}} tr.tot td{{font-weight:700;background:#fafafa}}
code{{background:#f0efe9;padding:1px 6px;border-radius:5px;font-size:12.5px}}
.pill{{font-size:11px;font-weight:600;padding:2px 9px;border-radius:999px;white-space:nowrap}}
.pill.ok{{background:#e6f4ea;color:#137333}} .pill.warn{{background:#fef7e0;color:#b06000}}
.pill.drop{{background:#fce8e6;color:#c5221f}} .pill.muted{{background:#eef0f2;color:#5f6368}}
.note{{background:#fff;border:1px solid var(--line);border-left:3px solid var(--yellow);
border-radius:8px;padding:14px 16px;margin:14px 0}}
.legend{{display:flex;gap:10px;flex-wrap:wrap;margin:10px 0 0}}
</style></head>
<body><div class="wrap">
<h1>Excel → UI coverage map</h1>
<p class="sub">How <code>Data from Miro.xlsx</code> maps onto the v3 Strategy Map data contract — what's covered, what's supplied by config, and what's dropped pending the Figma redesign.</p>

<div class="cards">
  <div class="card"><b>2</b><span>Excel sheets</span></div>
  <div class="card"><b>{sum(len(v) for sm in excel.values() for v in sm.values())}</b><span>JTBD rows (sheet 1)</span></div>
  <div class="card"><b>{len(enabler_rows)}</b><span>Enablers (sheet 2)</span></div>
  <div class="card"><b>{len(excel_phases)} / {sum(len(sm) for sm in excel.values())}</b><span>Phases / subphases</span></div>
</div>
<div class="legend">
  <span class="pill ok">mapped / derived</span>
  <span class="pill warn">from config</span>
  <span class="pill drop">dropped → Figma</span>
  <span class="pill muted">unused</span>
</div>

<h2>1 · Structure — current v3 vs incoming Excel</h2>
{struct_block()}
<div class="note"><strong>Structural mismatch.</strong> Phase names align 4-for-4
(Excel “Analyze” → v3 “Analyse”), but the subphase breakdown differs: v3 has placeholder
subphases that don't match the Excel's real ones (e.g. the “13-Slide Deck” is Excel
<code>2.6</code> under <em>Create</em>, not a Discover subphase). Per your call, the v3 app
shell is kept and the Excel content is mapped into it — adopting the Excel's subphases as
the real structure inside each phase.</div>

<h2>2 · Field coverage</h2>
{coverage_html}
<p class="muted">* <code>enablers</code> are sourced from sheet 2 and inserted within each
subphase (see §3). Association is currently by <em>Primary Phase</em> only — a subphase-level
tag would make placement exact (open question).</p>

<h2>3 · Enabler catalog (sheet 2) → inserted within subphases</h2>
<p class="muted">52 Tools · 29 Agents · 3 Super-Agents, grouped by the phase they belong to.</p>
{enabler_block()}
<div class="note"><strong>Open question.</strong> Sheet 2 ties each enabler to a
<em>Primary Phase</em> and a <em>Span</em>, but not to a specific subphase. To place them
inside the right subphase we either (a) show all of a phase's enablers on every subphase of
that phase, or (b) add a subphase column to the sheet. Flagging for your decision.</div>

</div></body></html>
"""

OUT.write_text(HTML, encoding="utf-8")
print(f"Wrote {OUT}  ({len(HTML):,} bytes)")
print(f"Excel: {sum(len(v) for sm in excel.values() for v in sm.values())} JTBD, "
      f"{len(enabler_rows)} enablers, {len(excel_phases)} phases")
